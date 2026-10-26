"""Load the capability catalog from YAML, JSON, or an Excel role mapping.

The default catalog ships with the package. Users may upload their own file
for a session; it is parsed in memory and never written to disk.
"""

from __future__ import annotations

import io
import json
import re
from functools import lru_cache
from importlib import resources
from typing import Any

import yaml
from pydantic import BaseModel, Field

from talent_estimation.domain.models import AiCategory, Scope


class RoleProfile(BaseModel):
    """Description of one role in the catalog."""

    code: str
    name: str
    team: str = "data"
    aliases: list[str] = Field(default_factory=list)
    summary: str = ""
    capability_areas: dict[str, str] = Field(default_factory=dict)
    in_scope: str = ""
    out_of_scope: str = ""
    primary_keywords: str = ""
    usually_supporting_on: str = ""
    substitutes: str = ""
    pillars: dict[str, list[str]] = Field(default_factory=dict)

    def describe(self) -> str:
        """Return a compact text description used in prompts.

        Returns
        -------
        str
            Multi-line description of the role.
        """
        pillars = "; ".join(f"{p}: {', '.join(s)}" for p, s in self.pillars.items())
        areas = "\n".join(f"    - {k}: {v}" for k, v in self.capability_areas.items())
        lines = [f"{self.code} - {self.name}"]
        if self.summary:
            lines.append(f"  Ringkasan: {self.summary}")
        if areas:
            lines.append(f"  Area capability:\n{areas}")
        lines += [
            f"  Lingkup utama: {self.in_scope}",
            f"  Di luar lingkup: {self.out_of_scope}",
            f"  Kata kunci pekerjaan utama: {self.primary_keywords}",
            f"  Biasanya pendukung pada: {self.usually_supporting_on}",
        ]
        if self.substitutes:
            lines.append(f"  Substitusi: {self.substitutes}")
        lines.append(f"  Pilar skill: {pillars}")
        return "\n".join(lines)


class Rule(BaseModel):
    """A decision rule followed by the agents."""

    id: str
    title: str
    text: str


class Catalog(BaseModel):
    """The full capability catalog."""

    version: str = "1.0"
    teams: dict[str, dict[str, Any]] = Field(default_factory=dict)
    roles: dict[str, RoleProfile]
    competency_scale: dict[int, str] = Field(default_factory=dict)
    rules: list[Rule] = Field(default_factory=list)
    allocation_guide: list[dict[str, str]] = Field(default_factory=list)
    classification_rubric: dict[str, str] = Field(default_factory=dict)
    actor_rubric: dict[str, str] = Field(default_factory=dict)
    ai_categories: list[AiCategory] = Field(default_factory=list)

    def roles_for_scope(self, scope: Scope) -> list[str]:
        """Return the role codes that belong to a scope.

        Parameters
        ----------
        scope : {"all", "data", "design"}
            Team scope selected by the user.

        Returns
        -------
        list of str
            Role codes in catalog order.
        """
        if scope == "all":
            return list(self.roles)
        return [code for code, role in self.roles.items() if role.team == scope]

    def ai_category(self, key: str) -> AiCategory | None:
        """Return the AI category with the given key, if it exists."""
        return next((c for c in self.ai_categories if c.key == key), None)

    def match_role(self, text: str) -> list[str]:
        """Find catalog roles mentioned in a free-text role hint.

        Parameters
        ----------
        text : str
            Role column of a WBS row, for example ``"Data Engineer, BA"``.

        Returns
        -------
        list of str
            Matching role codes. Short aliases match whole words only.
        """
        found: list[str] = []
        lowered = text.lower()
        for code, role in self.roles.items():
            for alias in [role.name.lower(), *role.aliases]:
                pattern = rf"\b{re.escape(alias.lower())}\d*\b"
                if re.search(pattern, lowered):
                    found.append(code)
                    break
        return found


def _from_mapping(data: dict[str, Any]) -> Catalog:
    roles = {code: RoleProfile(code=code, **body) for code, body in data.get("roles", {}).items()}
    return Catalog(**{**data, "roles": roles})


@lru_cache(maxsize=1)
def load_default_catalog() -> Catalog:
    """Load the catalog that ships with the package.

    Returns
    -------
    Catalog
        The default catalog.
    """
    text = (
        resources.files("talent_estimation.catalog")
        .joinpath("default_catalog.yaml")
        .read_text(encoding="utf-8")
    )
    return _from_mapping(yaml.safe_load(text))


def _split_skills(value: str) -> list[str]:
    return [s.strip() for s in re.split(r"[·;\n]", value) if s.strip()]


def _catalog_from_excel(content: bytes) -> Catalog:
    """Read a role mapping workbook with ``Scope_Capability`` and ``Pilar_Skill`` sheets."""
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    base = load_default_catalog().model_copy(deep=True)
    roles: dict[str, RoleProfile] = {}

    if "Scope_Capability" in workbook.sheetnames:
        rows = list(workbook["Scope_Capability"].iter_rows(values_only=True))
        header_index = next(
            (
                i
                for i, row in enumerate(rows)
                if row and str(row[0] or "").strip().lower() == "kode"
            ),
            None,
        )
        if header_index is not None:
            for row in rows[header_index + 1 :]:
                if not row or not row[0] or not row[1]:
                    continue
                code = str(row[0]).strip()
                if len(code) > 6:
                    continue
                default = base.roles.get(code)
                roles[code] = RoleProfile(
                    code=code,
                    name=str(row[1]).strip(),
                    team=default.team if default else "data",
                    aliases=default.aliases if default else [str(row[1]).strip().lower()],
                    in_scope=str(row[2] or ""),
                    out_of_scope=str(row[3] or ""),
                    primary_keywords=str(row[4] or ""),
                    usually_supporting_on=str(row[5] or ""),
                    pillars={},
                )

    if "Pilar_Skill" in workbook.sheetnames:
        for row in workbook["Pilar_Skill"].iter_rows(values_only=True):
            if not row or len(row) < 3 or not row[0] or not row[1]:
                continue
            code = str(row[0]).strip()
            if code in roles and row[2]:
                roles[code].pillars[str(row[1]).strip()] = _split_skills(str(row[2]))

    for code, role in roles.items():
        if not role.pillars and code in base.roles:
            role.pillars = base.roles[code].pillars

    if not roles:
        raise ValueError(
            "Workbook tidak memuat sheet Scope_Capability dengan kolom Kode dan Capability."
        )
    base.roles = roles
    return base


def load_catalog_file(filename: str, content: bytes) -> Catalog:
    """Parse an uploaded catalog file.

    Parameters
    ----------
    filename : str
        Original file name, used to detect the format.
    content : bytes
        File content. It is parsed in memory only.

    Returns
    -------
    Catalog
        The parsed catalog.

    Raises
    ------
    ValueError
        If the format is not supported or the content is invalid.
    """
    name = filename.lower()
    if name.endswith((".yaml", ".yml")):
        return _from_mapping(yaml.safe_load(content.decode("utf-8")))
    if name.endswith(".json"):
        return _from_mapping(json.loads(content.decode("utf-8")))
    if name.endswith(".xlsx"):
        return _catalog_from_excel(content)
    raise ValueError("Format katalog harus .yaml, .json, atau .xlsx.")
