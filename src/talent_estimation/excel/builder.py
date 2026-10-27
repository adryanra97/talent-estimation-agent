"""Build the validation workbook from a run.

The layout follows the original validation template: yellow cells hold inputs
written by the agents or the reviewer, grey cells hold formulas, and green
cells hold key results. Excel recalculates every formula when the file is
opened, so a reviewer can change an input and see all totals update.
"""

from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter as col_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from talent_estimation import __version__
from talent_estimation.catalog.loader import Catalog
from talent_estimation.domain.models import CLASSIFICATIONS, EF_FACTORS, TCF_FACTORS, RunResult
from talent_estimation.excel.styles import (
    BORDER,
    CALC_FILL,
    CENTER,
    F_BOLD,
    F_HDR,
    F_IN,
    F_NORM,
    F_NOTE,
    HDR_FILL,
    IN_FILL,
    RES_FILL,
    SEC_FILL,
    WARN_FILL,
    WRAP,
    WRAPC,
    as_calc,
    as_input,
    header,
    page,
    put,
    section,
    title,
    widths,
)

N1 = 7
"""First data row of the Breakdown_WBS sheet."""

FINDING_LABELS = {
    "gap": "Gap WBS",
    "overlap": "Tumpang tindih",
    "outside_wbs": "Di luar WBS",
    "divergence": "Selisih metode",
    "data_quality": "Kualitas data",
    "assumption": "Asumsi",
    "reclassified": "Klasifikasi ulang",
}


class WorkbookBuilder:
    """Create the validation workbook for one run.

    Parameters
    ----------
    run : RunResult
        A recalculated run.
    catalog : Catalog
        Capability catalog used in the run.
    """

    def __init__(self, run: RunResult, catalog: Catalog) -> None:
        self.run = run
        self.catalog = catalog
        self.roles = run.roles
        self.n_rows = max(150, len(run.breakdown) + 30)
        self.last_row = N1 + self.n_rows - 1
        self.n_months = max(1, len(run.monthly))
        self.wb = Workbook()
        self.fase_rows: dict[tuple[str, str], int] = {}

    # ------------------------------------------------------------------ helpers
    def _name(self, code: str) -> str:
        role = self.catalog.roles.get(code)
        return role.name if role else code

    def _rng(self, col: str) -> str:
        return f"Breakdown_WBS!${col}${N1}:${col}${self.last_row}"

    # --------------------------------------------------------------- parameter
    def _parameter(self) -> None:
        ws = self.wb.create_sheet("Parameter")
        run, p = self.run, self.run.parameters
        title(
            ws,
            "PARAMETER PROYEK & ASUMSI HITUNG",
            "Sel kuning bertulisan biru adalah input. Sel abu-abu adalah rumus. "
            "Setiap asumsi punya keterangan sumber.",
            4,
        )
        widths(ws, {"A": 42, "B": 20, "C": 14, "D": 80})
        section(ws, 4, "A. Informasi proyek", 4)
        info = [
            (5, "Nama proyek", run.project.name, "Teks bebas"),
            (
                6,
                "Referensi dokumen permintaan",
                run.project.reference,
                "Nomor dokumen atau tanggal notulen",
            ),
            (7, "Tanggal analisis", date.today(), "Tanggal file ini dibuat"),
            (
                8,
                "Disusun oleh",
                f"Talent Estimation Agent v{__version__}",
                "Hasil agent, perlu review",
            ),
        ]
        for row, label, value, note in info:
            put(ws, f"A{row}", label, border=BORDER)
            as_input(put(ws, f"B{row}", value))
            put(ws, f"D{row}", note, F_NOTE, align=WRAP)
        ws["B7"].number_format = "dd-mmm-yyyy"

        section(ws, 10, "B. Periode analisis (pembagi FTE)", 4)
        put(ws, "A11", "Tanggal mulai", border=BORDER)
        as_input(put(ws, "B11", run.project.start), "dd-mmm-yyyy")
        put(ws, "A12", "Tanggal selesai", border=BORDER)
        as_input(put(ws, "B12", run.project.end), "dd-mmm-yyyy")
        put(ws, "A13", "Durasi override (bulan), opsional", border=BORDER)
        as_input(put(ws, "B13", p.duration_override_months), "0.0")
        put(ws, "D13", "Isi hanya bila periode bukan kelipatan bulan penuh.", F_NOTE)
        put(ws, "A14", "Durasi periode", F_BOLD, border=BORDER)
        as_calc(
            put(
                ws,
                "B14",
                '=IF(B13<>"",B13,IF(OR(B11="",B12=""),0,'
                "(YEAR(B12)-YEAR(B11))*12+MONTH(B12)-MONTH(B11)+1))",
            ),
            "0.0",
            bold=True,
        )
        put(ws, "C14", "bulan")
        put(ws, "D14", "Selisih bulan inklusif. FTE = person-months / durasi ini.", F_NOTE)
        put(ws, "A15", "Hari kerja periode (tanpa libur)", border=BORDER)
        as_calc(put(ws, "B15", '=IF(OR(B11="",B12=""),"",NETWORKDAYS(B11,B12,Libur))'), "0")

        section(ws, 17, "C. Parameter hitung", 4)
        params = [
            (18, "Hari kerja per bulan", p.days_per_month, "hari/bulan", "0"),
            (19, "Jam kerja per hari", p.hours_per_day, "jam/hari", "0"),
            (20, "Person-Hours Multiplier (PHM)", p.phm, "jam/UCP", "0.00"),
            (21, "Faktor peran Pendukung (UCP)", p.support_factor, "x bobot", "0.00"),
            (22, "Toleransi selaras", p.tolerance, "+/-", "0%"),
        ]
        notes = {
            18: "Asumsi hari kerja efektif per bulan.",
            19: "Standar jam kerja harian.",
            20: "Jam kerja per poin UCP. Literatur 15-20 untuk use case software; "
            "item WBS lebih kecil sehingga nilai bawaan dikalibrasi ke 9.",
            21: "Item di mana role hanya mendukung dihitung sebagian dari bobotnya.",
            22: "Hasil dianggap selaras bila selisih dengan permintaan di bawah nilai ini.",
        }
        for row, label, value, unit, fmt in params:
            put(ws, f"A{row}", label, border=BORDER)
            as_input(put(ws, f"B{row}", value), fmt)
            put(ws, f"C{row}", unit)
            put(ws, f"D{row}", notes[row], F_NOTE, align=WRAP)

        section(ws, 24, "D. Bobot UCP (Use Case Points, Karner 1993)", 4)
        header(ws, 25, ["Kategori", "Bobot aktor", "Bobot use case", "Keterangan"])
        for i, cat in enumerate(CLASSIFICATIONS):
            row = 26 + i
            put(ws, f"A{row}", cat, F_BOLD, border=BORDER)
            as_input(put(ws, f"B{row}", p.actor_weights[cat]), "0")
            as_input(put(ws, f"C{row}", p.use_case_weights[cat]), "0")
            put(ws, f"D{row}", self.catalog.classification_rubric.get(cat, ""), F_NOTE)

        section(ws, 30, "E. Hari libur dalam periode", 4)
        header(ws, 31, ["Tanggal", "", "", "Keterangan"])
        for i in range(30):
            row = 32 + i
            value = p.holidays[i] if i < len(p.holidays) else None
            as_input(put(ws, f"A{row}", value), "dd-mmm-yyyy")
        names = {
            "Mulai": "$B$11",
            "Selesai": "$B$12",
            "Durasi_Bulan": "$B$14",
            "Hari_per_Bulan": "$B$18",
            "Jam_per_Hari": "$B$19",
            "PHM": "$B$20",
            "Faktor_Pendukung": "$B$21",
            "Toleransi": "$B$22",
            "Libur": "$A$32:$A$61",
        }
        for name, ref in names.items():
            self.wb.defined_names[name] = DefinedName(name, attr_text=f"Parameter!{ref}")
        ws.freeze_panes = "A4"
        page(ws)

    # --------------------------------------------------------------- breakdown
    def _breakdown(self) -> None:
        ws = self.wb.create_sheet("Breakdown_WBS")
        last = self.last_row
        title(
            ws,
            "BREAKDOWN WBS KE ROLE (SUMBER SEMUA ANGKA)",
            "Satu baris = satu item WBS level terbawah x satu role. Kolom kuning diisi agent "
            "dan direview. Kolom abu-abu rumus. Klasifikasi dan alokasi wajib beralasan.",
            24,
        )
        put(ws, "A3", "Baris terisi:", F_BOLD)
        as_calc(put(ws, "C3", f"=COUNTA(B{N1}:B{last})"), "0")
        put(ws, "D3", "Perlu review:", F_BOLD)
        as_calc(put(ws, "E3", f'=COUNTIF(R{N1}:R{last},"Ya")'), "0")
        put(ws, "F3", "Baris belum lengkap:", F_BOLD)
        as_calc(
            put(
                ws,
                "H3",
                f'=SUMPRODUCT((B{N1}:B{last}<>"")*(((J{N1}:J{last}="")+(K{N1}:K{last}="")'
                f'+(N{N1}:N{last}="")+(P{N1}:P{last}="")+(F{N1}:F{last}="")'
                f'+(G{N1}:G{last}="")+(H{N1}:H{last}=""))>0))',
            ),
            "0",
        )
        put(ws, "I3", "Harus 0 sebelum hasil dipakai", F_NOTE)
        for first, last_col, text in (
            ("B", "I", "DARI WBS"),
            ("J", "R", "PENILAIAN AGENT (direview manusia)"),
            ("S", "W", "HASIL HITUNG (rumus)"),
        ):
            ws.merge_cells(f"{first}5:{last_col}5")
            for col in range(ws[f"{first}5"].column, ws[f"{last_col}5"].column + 1):
                ws.cell(row=5, column=col).fill = SEC_FILL
            put(ws, f"{first}5", text, F_BOLD, SEC_FILL, CENTER)
        heads = [
            "No", "WBS ID", "Item WBS", "Deskripsi", "Role di WBS (petunjuk)", "Start", "End",
            "Hari kerja item", "Kelompok", "Role", "Peran", "Qty", "Jumlah langkah",
            "Klasifikasi UCP", "Alasan klasifikasi & peran", "Alokasi FTE (bottom-up)",
            "Alasan alokasi FTE", "Perlu review?", "Bobot dasar UCP", "Bobot efektif UCP",
            "Person-days (bottom-up)", "Jam UCP teralokasi", "Person-days UCP teralokasi", "Sumber",
        ]  # fmt: skip
        header(ws, 6, heads)
        ws.row_dimensions[6].height = 42
        widths(
            ws,
            {
                "A": 5, "B": 10, "C": 30, "D": 40, "E": 18, "F": 11, "G": 11, "H": 8, "I": 16,
                "J": 8, "K": 11, "L": 6, "M": 9, "N": 12, "O": 46, "P": 11, "Q": 40, "R": 9,
                "S": 9, "T": 10, "U": 12, "V": 11, "W": 12, "X": 12,
            },
        )  # fmt: skip
        n_roles = len(self.roles)
        r_first, r_last = 7, 6 + n_roles
        for row in range(N1, last + 1):
            as_calc(ws.cell(row=row, column=1, value=f'=IF(B{row}="","",ROW()-{N1 - 1})'), "0")
            for col in range(2, 19):
                cell = ws.cell(row=row, column=col)
                as_input(cell)
                if col in (6, 7):
                    cell.number_format = "dd-mmm-yy"
                if col in (3, 4, 15, 17):
                    cell.alignment = WRAP
                if col == 16:
                    cell.number_format = "0.00"
            as_input(ws.cell(row=row, column=24))
            as_calc(
                ws.cell(
                    row=row,
                    column=19,
                    value=f'=IF(OR(B{row}="",N{row}=""),"",'
                    f"INDEX(Parameter!$C$26:$C$28,MATCH(N{row},Parameter!$A$26:$A$28,0)))",
                ),
                "0",
            )
            as_calc(
                ws.cell(
                    row=row,
                    column=20,
                    value=f'=IF(S{row}="","",S{row}*L{row}*IF(K{row}="Pendukung",Faktor_Pendukung,1))',
                ),
                "0.00",
            )
            as_calc(
                ws.cell(
                    row=row,
                    column=21,
                    value=f'=IF(OR(B{row}="",P{row}="",H{row}=""),"",H{row}*P{row}*L{row})',
                ),
                "0.0",
            )
            as_calc(
                ws.cell(
                    row=row,
                    column=22,
                    value=(
                        f'=IF(OR(T{row}="",J{row}=""),"",'
                        f"IF(SUMIFS($T${N1}:$T${last},$J${N1}:$J${last},J{row})=0,0,"
                        f"T{row}/SUMIFS($T${N1}:$T${last},$J${N1}:$J${last},J{row})"
                        f"*INDEX(Ringkasan!$F${r_first}:$F${r_last},"
                        f"MATCH(J{row},Ringkasan!$A${r_first}:$A${r_last},0))))"
                    ),
                ),
                "#,##0.0",
            )
            as_calc(
                ws.cell(row=row, column=23, value=f'=IF(V{row}="","",V{row}/Jam_per_Hari)'), "0.0"
            )
        role_list = ",".join(self.roles)
        for rng, allowed in (
            (f"J{N1}:J{last}", f'"{role_list}"'),
            (f"K{N1}:K{last}", '"Utama,Pendukung"'),
            (f"N{N1}:N{last}", '"Simple,Average,Complex"'),
            (f"R{N1}:R{last}", '"Ya,Tidak"'),
        ):
            dv = DataValidation(
                type="list", formula1=allowed, allow_blank=True, showErrorMessage=True
            )
            ws.add_data_validation(dv)
            dv.add(rng)
        dv = DataValidation(
            type="decimal", operator="between", formula1="0", formula2="5", allow_blank=True,
            showErrorMessage=True, errorTitle="Alokasi FTE", error="Isi angka 0-5.",
        )  # fmt: skip
        ws.add_data_validation(dv)
        dv.add(f"P{N1}:P{last}")
        ws.conditional_formatting.add(
            f"R{N1}:R{last}", CellIsRule(operator="equal", formula=['"Ya"'], fill=WARN_FILL)
        )
        for i, r in enumerate(self.run.breakdown):
            row = N1 + i
            values = [
                r.item_id, r.item_name, r.description, r.role_hint, r.start, r.end, r.workdays,
                r.group, r.role, r.involvement, r.quantity, r.steps, r.classification,
                r.classification_reason, r.allocation_fte, r.allocation_reason,
                "Ya" if r.needs_review else "Tidak",
            ]  # fmt: skip
            for j, value in enumerate(values):
                ws.cell(row=row, column=2 + j, value=value)
            ws.cell(row=row, column=24, value="WBS" if r.source == "wbs" else "Dokumen tambahan")
        ws.freeze_panes = f"D{N1}"
        ws.auto_filter.ref = f"A6:X{last}"
        page(ws)

    # --------------------------------------------------------------------- UCP
    def _ucp(self, code: str) -> None:
        ws = self.wb.create_sheet(f"UCP_{code}")
        factors = next((f for f in self.run.factors if f.role == code), None)
        last = self.last_row
        title(
            ws,
            f"UCP: {self._name(code)} ({code})",
            "Pendekatan 1, Use Case Points. Jumlah item dan klasifikasi diambil otomatis dari "
            "Breakdown_WBS. Aktor dan rating TCF/EF diisi dengan alasan tertulis.",
            7,
        )
        put(ws, "A3", "Kode:", F_BOLD)
        put(ws, "B3", code, F_BOLD)
        widths(ws, {"A": 9, "B": 54, "C": 12, "D": 12, "E": 12, "F": 14, "G": 64})
        section(
            ws, 5, "A. UNADJUSTED ACTOR WEIGHT (UAW): pihak di luar role ini yang berinteraksi", 7
        )
        header(ws, 6, ["No", "Aktor", "Kategori", "Jumlah", "Bobot", "Hasil", "Sumber / alasan"])
        actors = factors.actors if factors else []
        for i in range(12):
            row = 7 + i
            actor = actors[i] if i < len(actors) else None
            put(ws, f"A{row}", i + 1, align=CENTER, border=BORDER)
            as_input(put(ws, f"B{row}", actor.name if actor else None))
            as_input(put(ws, f"C{row}", actor.category if actor else None))
            as_input(put(ws, f"D{row}", actor.count if actor else 1), "0")
            as_calc(
                put(
                    ws,
                    f"E{row}",
                    f'=IF(C{row}="","",INDEX(Parameter!$B$26:$B$28,'
                    f"MATCH(C{row},Parameter!$A$26:$A$28,0)))",
                ),
                "0",
            )
            as_calc(put(ws, f"F{row}", f'=IF(C{row}="","",E{row}*D{row})'), "0.0")
            as_input(put(ws, f"G{row}", actor.source if actor else None))
            ws[f"G{row}"].alignment = WRAP
        dv = DataValidation(type="list", formula1='"Simple,Average,Complex"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("C7:C18")
        put(ws, "B19", "Total UAW", F_BOLD, border=BORDER)
        as_calc(put(ws, "F19", "=SUM(F7:F18)"), "0.0", RES_FILL, True)

        section(ws, 21, "B. UNADJUSTED USE CASE WEIGHT (UUCW): otomatis dari Breakdown_WBS", 7)
        header(
            ws,
            22,
            ["Kategori", "Rubrik", "Bobot", "Jumlah Utama", "Jumlah Pendukung", "Hasil", "Catatan"],
        )
        for i, (cat, rubric) in enumerate(
            (("Simple", "1-3 langkah"), ("Average", "4-7 langkah"), ("Complex", ">7 langkah"))
        ):
            row = 23 + i
            put(ws, f"A{row}", cat, F_BOLD, border=BORDER)
            put(ws, f"B{row}", rubric, border=BORDER)
            as_calc(
                put(
                    ws,
                    f"C{row}",
                    f"=INDEX(Parameter!$C$26:$C$28,MATCH(A{row},Parameter!$A$26:$A$28,0))",
                ),
                "0",
            )
            for col, involvement in (("D", "Utama"), ("E", "Pendukung")):
                as_calc(
                    put(
                        ws,
                        f"{col}{row}",
                        f"=SUMIFS({self._rng('L')},{self._rng('J')},$B$3,"
                        f'{self._rng("N")},$A{row},{self._rng("K")},"{involvement}")',
                    ),
                    "0.0",
                )
            as_calc(put(ws, f"F{row}", f"=C{row}*(D{row}+E{row}*Faktor_Pendukung)"), "0.0")
        put(ws, "G23", "Hasil = bobot x (Utama + Pendukung x Faktor_Pendukung)", F_NOTE, align=WRAP)
        put(ws, "B26", "Total UUCW", F_BOLD, border=BORDER)
        as_calc(put(ws, "F26", "=SUM(F23:F25)"), "0.0", RES_FILL, True)
        put(ws, "B27", "Cek terhadap bobot efektif di Breakdown_WBS", F_NOTE)
        as_calc(
            put(
                ws,
                "F27",
                f"=IF(ABS(F26-SUMIFS(Breakdown_WBS!$T${N1}:$T${last},"
                f'Breakdown_WBS!$J${N1}:$J${last},$B$3))<0.001,"OK","SELISIH")',
            )
        )
        put(ws, "B29", "TOTAL UUCP = UAW + UUCW", F_BOLD, border=BORDER)
        as_calc(put(ws, "F29", "=F19+F26"), "0.0", RES_FILL, True)

        def factor_block(start: int, rows: tuple[tuple[str, str, float], ...]) -> None:
            for i, (fcode, fname, weight) in enumerate(rows):
                row = start + i
                rating = factors.rating(fcode) if factors else 3
                evidence = factors.evidence.get(fcode, "") if factors else ""
                put(ws, f"A{row}", fcode, F_BOLD, border=BORDER, align=CENTER)
                put(ws, f"B{row}", fname, border=BORDER)
                put(ws, f"C{row}", weight, border=BORDER, fmt="0.0")
                as_input(put(ws, f"D{row}", rating), "0")
                as_calc(put(ws, f"E{row}", f"=C{row}*D{row}"), "0.0")
                as_calc(
                    put(
                        ws,
                        f"F{row}",
                        f'=IF(D{row}="","",IF(D{row}=3,"Default",'
                        f'IF(G{row}="","BUKTI KOSONG","Diubah")))',
                    )
                )
                as_input(put(ws, f"G{row}", evidence or None))
                ws[f"G{row}"].alignment = WRAP

        section(
            ws,
            31,
            "C. TECHNICAL COMPLEXITY FACTOR (TCF). Rating 0-5, default 3; ubah hanya dengan bukti",
            7,
        )
        header(
            ws,
            32,
            ["Kode", "Faktor teknis", "Bobot", "Rating (0-5)", "Nilai", "Status", "Alasan / bukti"],
        )
        factor_block(33, TCF_FACTORS)
        put(ws, "B46", "Total Technical Factor (TFactor)", F_BOLD, border=BORDER)
        as_calc(put(ws, "E46", "=SUM(E33:E45)"), "0.0", bold=True)
        put(ws, "B47", "TCF = 0,6 + (0,01 x TFactor)", F_BOLD, border=BORDER)
        as_calc(put(ws, "E47", "=0.6+(0.01*E46)"), "0.000", RES_FILL, True)
        section(
            ws,
            49,
            "D. ENVIRONMENTAL FACTOR (EF). E1-E6 tinggi menguntungkan, E7-E8 tinggi memberatkan",
            7,
        )
        header(
            ws,
            50,
            [
                "Kode",
                "Faktor lingkungan",
                "Bobot",
                "Rating (0-5)",
                "Nilai",
                "Status",
                "Alasan / bukti",
            ],
        )
        factor_block(51, EF_FACTORS)
        put(ws, "B59", "Total Environmental Factor (EFactor)", F_BOLD, border=BORDER)
        as_calc(put(ws, "E59", "=SUM(E51:E58)"), "0.0", bold=True)
        put(ws, "B60", "EF = 1,4 + (-0,03 x EFactor)", F_BOLD, border=BORDER)
        as_calc(put(ws, "E60", "=1.4+(-0.03*E59)"), "0.000", RES_FILL, True)
        for rng in ("D33:D45", "D51:D58"):
            dv = DataValidation(
                type="whole", operator="between", formula1="0", formula2="5", allow_blank=True
            )
            ws.add_data_validation(dv)
            dv.add(rng)
        ws.conditional_formatting.add(
            "F33:F58",
            CellIsRule(
                operator="equal",
                formula=['"BUKTI KOSONG"'],
                fill=PatternFill("solid", fgColor="F8CBAD"),
            ),
        )

        section(ws, 62, "E. HASIL: UCP, effort, dan FTE", 7)
        results = [
            (63, "Total UUCP", "=F29", "0.0", "UAW + UUCW"),
            (64, "TCF", "=E47", "0.000", "Bagian C"),
            (65, "EF", "=E60", "0.000", "Bagian D"),
            (66, "Adjusted UCP = UUCP x TCF x EF", "=F63*F64*F65", "0.0", "Rumus baku UCP"),
            (67, "Person-Hours Multiplier (PHM)", "=PHM", "0.00", "Parameter"),
            (68, "Man-hours = UCP x PHM", "=F66*F67", "#,##0", "Total jam kerja role ini"),
            (
                69,
                "Person-days = man-hours / jam per hari",
                "=F68/Jam_per_Hari",
                "#,##0.0",
                "Parameter",
            ),
            (
                70,
                "Person-months = person-days / hari per bulan",
                "=F69/Hari_per_Bulan",
                "0.00",
                "Parameter",
            ),
            (71, "Durasi periode (bulan)", "=Durasi_Bulan", "0.0", "Parameter"),
            (
                72,
                "FTE rata-rata = person-months / durasi",
                "=IF(F71=0,0,F70/F71)",
                "0.00",
                "Dibawa ke Ringkasan",
            ),
        ]
        for row, label, formula, fmt, note in results:
            key = row in (66, 68, 72)
            put(ws, f"B{row}", label, F_BOLD if key else F_NORM, border=BORDER)
            as_calc(put(ws, f"F{row}", formula), fmt, RES_FILL if key else CALC_FILL, key)
            put(ws, f"G{row}", note, F_NOTE)
        ws.freeze_panes = "A5"
        page(ws, landscape=False)

    # --------------------------------------------------------------- bottom-up
    def _bottom_up(self) -> None:
        ws = self.wb.create_sheet("BottomUp")
        title(
            ws,
            "BOTTOM-UP: EFFORT DARI JADWAL WBS x ALOKASI FTE",
            "Pendekatan 2. Person-days item = hari kerja item x alokasi FTE x Qty. "
            "Alokasi diestimasi dari beban kerja, bukan menyalin jumlah orang yang diminta.",
            8,
        )
        widths(ws, {"A": 9, "B": 22, "C": 14, "D": 14, "E": 14, "F": 14, "G": 12, "H": 16})
        section(ws, 4, "A. Hasil bottom-up per role", 8)
        header(
            ws,
            5,
            [
                "Kode",
                "Role",
                "Person-days",
                "Man-hours",
                "Person-months",
                "FTE rata-rata",
                "Jumlah baris",
                "Rata-rata alokasi",
            ],
        )
        for i, code in enumerate(self.roles):
            row = 6 + i
            put(ws, f"A{row}", code, F_BOLD, border=BORDER)
            put(ws, f"B{row}", self._name(code), border=BORDER)
            as_calc(
                put(ws, f"C{row}", f"=SUMIFS({self._rng('U')},{self._rng('J')},$A{row})"), "#,##0.0"
            )
            as_calc(put(ws, f"D{row}", f"=C{row}*Jam_per_Hari"), "#,##0")
            as_calc(put(ws, f"E{row}", f"=C{row}/Hari_per_Bulan"), "0.00")
            as_calc(
                put(ws, f"F{row}", f"=IF(Durasi_Bulan=0,0,E{row}/Durasi_Bulan)"),
                "0.00",
                RES_FILL,
                True,
            )
            as_calc(
                put(ws, f"G{row}", f'=COUNTIFS({self._rng("J")},$A{row},{self._rng("P")},"<>")'),
                "0",
            )
            as_calc(
                put(
                    ws,
                    f"H{row}",
                    f"=IF(G{row}=0,0,SUMIFS({self._rng('P')},{self._rng('J')},$A{row})/G{row})",
                ),
                "0.00",
            )
        total = 6 + len(self.roles)
        put(ws, f"B{total}", "Total", F_BOLD, border=BORDER)
        for col, fmt in (
            ("C", "#,##0.0"),
            ("D", "#,##0"),
            ("E", "0.00"),
            ("F", "0.00"),
            ("G", "0"),
        ):
            as_calc(put(ws, f"{col}{total}", f"=SUM({col}6:{col}{total - 1})"), fmt, RES_FILL, True)
        start = total + 3
        section(ws, start, "B. Panduan alokasi FTE per item", 8)
        header(ws, start + 1, ["Peran", "Klasifikasi", "Rentang alokasi", "", "", "", "", "Makna"])
        for i, guide in enumerate(self.catalog.allocation_guide):
            row = start + 2 + i
            put(ws, f"A{row}", guide.get("involvement", ""), F_BOLD, border=BORDER)
            put(ws, f"B{row}", guide.get("classification", ""), border=BORDER)
            put(ws, f"C{row}", guide.get("range", ""), border=BORDER)
            put(ws, f"H{row}", guide.get("meaning", ""), F_NOTE)
        page(ws)

    # ---------------------------------------------------------- monthly phases
    def _fase(self) -> None:
        ws = self.wb.create_sheet("Fase_Bulanan")
        months = self.n_months
        roles = self.roles
        n = len(roles)
        peak_col = 3 + months
        title(
            ws,
            "FASE BULANAN: SEBARAN KEBUTUHAN PER BULAN (FTE) DAN PUNCAK",
            "Person-days tiap item disebar ke bulan sesuai irisan hari kerja Start-End item. "
            "FTE bulan = person-days bulan / hari kerja efektif bulan.",
            peak_col,
        )
        widths(ws, {"A": 11, "B": 12, **{col_letter(c): 10 for c in range(3, peak_col + 1)}})
        section(ws, 4, "A. FTE per bulan", peak_col)
        header(ws, 5, ["Kode", "Metode"] + [""] * months + ["Puncak"])
        for j in range(months):
            col = 3 + j
            formula = (
                '=IF(Mulai="","",DATE(YEAR(Mulai),MONTH(Mulai),1))'
                if j == 0
                else f'=IF({col_letter(col - 1)}5="","",EDATE({col_letter(col - 1)}5,1))'
            )
            cell = ws.cell(row=5, column=col, value=formula)
            cell.number_format = "mmm-yy"
            cell.font = F_HDR
            cell.fill = HDR_FILL
            cell.alignment = WRAPC
            cell.border = BORDER
        cap_row = 6 + 2 * (n + 1) + 1
        det0 = cap_row + 6
        detl = det0 + self.n_rows - 1
        ucp0 = 7
        bu0 = ucp0 + months
        row = 6
        for method, base in (("UCP", ucp0), ("Bottom-up", bu0)):
            first = row
            for code in roles:
                put(ws, f"A{row}", code, F_BOLD, border=BORDER)
                put(ws, f"B{row}", method, border=BORDER)
                for j in range(months):
                    mc = col_letter(3 + j)
                    dc = col_letter(base + j)
                    as_calc(
                        put(
                            ws,
                            f"{mc}{row}",
                            f'=IF(OR({mc}$5="",{mc}$5>Selesai),"",IF({mc}${cap_row}=0,0,'
                            f"SUMIFS({dc}${det0}:{dc}${detl},$B${det0}:$B${detl},$A{row})/{mc}${cap_row}))",
                        ),
                        "0.00",
                    )
                as_calc(
                    put(
                        ws,
                        f"{col_letter(peak_col)}{row}",
                        f"=MAX(C{row}:{col_letter(peak_col - 1)}{row})",
                    ),
                    "0.00",
                    RES_FILL,
                    True,
                )
                self.fase_rows[(method, code)] = row
                row += 1
            put(ws, f"A{row}", "Total", F_BOLD, border=BORDER)
            put(ws, f"B{row}", method, F_BOLD, border=BORDER)
            for j in range(months):
                mc = col_letter(3 + j)
                as_calc(
                    put(
                        ws,
                        f"{mc}{row}",
                        f'=IF(OR({mc}$5="",{mc}$5>Selesai),"",SUM({mc}{first}:{mc}{row - 1}))',
                    ),
                    "0.00",
                    RES_FILL,
                    True,
                )
            as_calc(
                put(
                    ws,
                    f"{col_letter(peak_col)}{row}",
                    f"=MAX(C{row}:{col_letter(peak_col - 1)}{row})",
                ),
                "0.00",
                RES_FILL,
                True,
            )
            self.fase_rows[(method, "TOTAL")] = row
            row += 1
        put(ws, f"A{cap_row}", "Hari kerja efektif", F_BOLD)
        for j in range(months):
            mc = col_letter(3 + j)
            as_calc(
                put(
                    ws,
                    f"{mc}{cap_row}",
                    f'=IF(OR({mc}$5="",Selesai=""),0,IF({mc}$5>Selesai,0,'
                    f"NETWORKDAYS(MAX({mc}$5,Mulai),MIN(EOMONTH({mc}$5,0),Selesai),Libur)))",
                ),
                "0",
            )
        last_ucp = col_letter(ucp0 + months - 1)
        last_bu = col_letter(bu0 + months - 1)
        put(ws, f"A{cap_row + 1}", "Cek sebaran UCP:", F_BOLD)
        as_calc(
            put(
                ws,
                f"C{cap_row + 1}",
                f"=IF(ABS(SUM({col_letter(ucp0)}{det0}:{last_ucp}{detl})-SUM(E{det0}:E{detl}))<0.01,"
                '"OK","Ada item di luar periode")',
            )
        )
        put(ws, f"A{cap_row + 2}", "Cek sebaran Bottom-up:", F_BOLD)
        as_calc(
            put(
                ws,
                f"C{cap_row + 2}",
                f"=IF(ABS(SUM({col_letter(bu0)}{det0}:{last_bu}{detl})-SUM(F{det0}:F{detl}))<0.01,"
                '"OK","Ada item di luar periode")',
            )
        )
        self.fase_checks = (f"C{cap_row + 1}", f"C{cap_row + 2}")
        section(ws, det0 - 2, "B. Detail sebaran person-days per item (otomatis)", 6)
        header(ws, det0 - 1, ["WBS ID", "Role", "Start", "End", "PD UCP", "PD Bottom-up"])
        for base in (ucp0, bu0):
            for j in range(months):
                cell = ws.cell(row=det0 - 1, column=base + j, value=f"={col_letter(3 + j)}$5")
                cell.number_format = "mmm-yy"
                cell.font = F_HDR
                cell.fill = HDR_FILL
                cell.alignment = WRAPC
                cell.border = BORDER
        put(ws, f"{col_letter(ucp0)}{det0 - 2}", "Person-days per bulan: UCP", F_BOLD, SEC_FILL)
        put(
            ws, f"{col_letter(bu0)}{det0 - 2}", "Person-days per bulan: Bottom-up", F_BOLD, SEC_FILL
        )
        for k in range(self.n_rows):
            r = det0 + k
            b = N1 + k
            as_calc(put(ws, f"A{r}", f'=IF(Breakdown_WBS!B{b}="","",Breakdown_WBS!B{b})'))
            as_calc(put(ws, f"B{r}", f'=IF(Breakdown_WBS!J{b}="","",Breakdown_WBS!J{b})'))
            as_calc(
                put(ws, f"C{r}", f'=IF(Breakdown_WBS!F{b}="","",Breakdown_WBS!F{b})'), "dd-mmm-yy"
            )
            as_calc(
                put(ws, f"D{r}", f'=IF(Breakdown_WBS!G{b}="","",Breakdown_WBS!G{b})'), "dd-mmm-yy"
            )
            as_calc(put(ws, f"E{r}", f'=IF(Breakdown_WBS!W{b}="","",Breakdown_WBS!W{b})'), "0.0")
            as_calc(put(ws, f"F{r}", f'=IF(Breakdown_WBS!U{b}="","",Breakdown_WBS!U{b})'), "0.0")
            for base, src in ((ucp0, "E"), (bu0, "F")):
                for j in range(months):
                    cl = col_letter(base + j)
                    mcol = f"{cl}${det0 - 1}"
                    formula = (
                        f'=IF(OR(${src}{r}="",$C{r}="",$D{r}="",{mcol}=""),"",${src}{r}*'
                        f"IF(NETWORKDAYS($C{r},$D{r},Libur)<=0,0,"
                        f"IF(MAX($C{r},{mcol})>MIN($D{r},EOMONTH({mcol},0)),0,"
                        f"NETWORKDAYS(MAX($C{r},{mcol}),MIN($D{r},EOMONTH({mcol},0)),Libur))"
                        f"/NETWORKDAYS($C{r},$D{r},Libur)))"
                    )
                    as_calc(put(ws, f"{cl}{r}", formula), "0.0")
        ws.freeze_panes = "C6"
        page(ws)

    # --------------------------------------------------------------- summary
    def _ringkasan(self) -> None:
        ws = self.wb.create_sheet("Ringkasan", 1)
        roles = self.roles
        title(
            ws,
            "RINGKASAN VALIDASI KEBUTUHAN TALENT",
            "Dua pendekatan independen (UCP dan Bottom-Up) dibandingkan dengan permintaan. "
            "Kebutuhan tambahan = kebutuhan total dikurangi talent existing.",
            17,
        )
        put(
            ws,
            "A3",
            '="Proyek: "&IF(Parameter!B5="","-",Parameter!B5)&"   |   Periode: "'
            '&IF(Parameter!B11="","?",TEXT(Parameter!B11,"dd-mmm-yyyy"))&" s.d. "'
            '&IF(Parameter!B12="","?",TEXT(Parameter!B12,"dd-mmm-yyyy"))'
            '&"   |   Durasi: "&TEXT(Durasi_Bulan,"0.0")&" bulan"',
            F_BOLD,
        )
        ws.merge_cells("A3:Q3")
        widths(
            ws,
            {
                "A": 8, "B": 16, "C": 11, "D": 11, "E": 12, "F": 11, "G": 11, "H": 11, "I": 11,
                "J": 11, "K": 11, "L": 11, "M": 11, "N": 11, "O": 13, "P": 13, "Q": 44,
            },
        )  # fmt: skip
        section(ws, 5, "1. Perbandingan kebutuhan per role (FTE = person-months / durasi)", 17)
        header(
            ws,
            6,
            [
                "Kode", "Role", "Diminta (FTE)", "Talent existing (FTE)", "Usulan pembanding (FTE)",
                "UCP man-hours", "UCP FTE rata-rata", "UCP FTE puncak", "Bottom-up man-hours",
                "Bottom-up FTE rata-rata", "Bottom-up FTE puncak", "Selisih UCP vs Bottom-up",
                "Tambahan UCP", "Tambahan Bottom-up", "UCP vs diminta", "Bottom-up vs diminta",
                "Kesimpulan",
            ],
        )  # fmt: skip
        ws.row_dimensions[6].height = 44
        inputs = {i.role: i for i in self.run.role_inputs}
        for i, code in enumerate(roles):
            row = 7 + i
            ri = inputs.get(code)
            put(ws, f"A{row}", code, F_BOLD, border=BORDER)
            put(ws, f"B{row}", self._name(code), border=BORDER)
            as_input(put(ws, f"C{row}", ri.requested_fte if ri else None), "0.00")
            as_input(put(ws, f"D{row}", ri.existing_fte if ri else 0), "0.00")
            as_input(put(ws, f"E{row}", ri.reference_fte if ri else None), "0.00")
            as_calc(put(ws, f"F{row}", f"=UCP_{code}!F68"), "#,##0")
            as_calc(put(ws, f"G{row}", f"=UCP_{code}!F72"), "0.00", RES_FILL, True)
            as_calc(
                put(ws, f"H{row}", f"=Fase_Bulanan!{self._peak()}{self.fase_rows[('UCP', code)]}"),
                "0.00",
            )
            as_calc(put(ws, f"I{row}", f"=BottomUp!D{6 + i}"), "#,##0")
            as_calc(put(ws, f"J{row}", f"=BottomUp!F{6 + i}"), "0.00", RES_FILL, True)
            as_calc(
                put(
                    ws,
                    f"K{row}",
                    f"=Fase_Bulanan!{self._peak()}{self.fase_rows[('Bottom-up', code)]}",
                ),
                "0.00",
            )
            as_calc(
                put(ws, f"L{row}", f'=IF(OR(J{row}=0,G{row}=0),"",(G{row}-J{row})/J{row})'),
                "+0%;-0%;0%",
            )
            as_calc(put(ws, f"M{row}", f"=MAX(G{row}-D{row},0)"), "0.00")
            as_calc(put(ws, f"N{row}", f"=MAX(J{row}-D{row},0)"), "0.00")
            for col, basis in (("O", "M"), ("P", "N")):
                as_calc(
                    put(
                        ws,
                        f"{col}{row}",
                        f'=IF(C{row}="","",IF({basis}{row}<C{row}*(1-Toleransi),"Lebih rendah",'
                        f'IF({basis}{row}>C{row}*(1+Toleransi),"Lebih tinggi","Selaras")))',
                    )
                )
            as_calc(
                put(
                    ws,
                    f"Q{row}",
                    f'=IF(O{row}="","Isi FTE yang diminta untuk melihat kesimpulan",'
                    f'IF(AND(O{row}="Lebih rendah",P{row}="Lebih rendah"),'
                    '"Kebutuhan tambahan lebih rendah dari permintaan (kedua metode)",'
                    f'IF(AND(O{row}="Lebih tinggi",P{row}="Lebih tinggi"),'
                    '"Kebutuhan tambahan lebih tinggi dari permintaan (kedua metode)",'
                    f'IF(AND(O{row}="Selaras",P{row}="Selaras"),"Selaras dengan permintaan",'
                    '"Dua metode tidak sepakat, perlu review manual"))))',
                )
            )
            ws[f"Q{row}"].alignment = WRAP
        total = 7 + len(roles)
        put(ws, f"B{total}", "Total", F_BOLD, border=BORDER)
        for col, fmt in (
            ("C", "0.00"), ("D", "0.00"), ("E", "0.00"), ("F", "#,##0"), ("G", "0.00"),
            ("I", "#,##0"), ("J", "0.00"), ("M", "0.00"), ("N", "0.00"),
        ):  # fmt: skip
            as_calc(put(ws, f"{col}{total}", f"=SUM({col}7:{col}{total - 1})"), fmt, RES_FILL, True)
        as_calc(
            put(ws, f"H{total}", f"=Fase_Bulanan!{self._peak()}{self.fase_rows[('UCP', 'TOTAL')]}"),
            "0.00",
            RES_FILL,
            True,
        )
        as_calc(
            put(
                ws,
                f"K{total}",
                f"=Fase_Bulanan!{self._peak()}{self.fase_rows[('Bottom-up', 'TOTAL')]}",
            ),
            "0.00",
            RES_FILL,
            True,
        )
        as_calc(
            put(ws, f"L{total}", f'=IF(OR(J{total}=0,G{total}=0),"",(G{total}-J{total})/J{total})'),
            "+0%;-0%;0%",
        )
        notes = [
            "Diminta = FTE yang diminta pada dokumen permintaan. Talent existing default 0. "
            "Kebutuhan tambahan (kolom M dan N) yang dibandingkan dengan permintaan.",
            "FTE puncak = bulan dengan kebutuhan tertinggi (lihat Fase_Bulanan). Total puncak bukan "
            "penjumlahan puncak per role.",
            "Penentuan nama talent berada di luar cakupan file ini.",
        ]
        for i, text in enumerate(notes):
            row = total + 1 + i
            ws.merge_cells(f"A{row}:Q{row}")
            put(ws, f"A{row}", text, F_NOTE, align=WRAP)
            ws.row_dimensions[row].height = 24
        start = total + 5
        section(ws, start, "2. FTE per bulan (6 bulan pertama; selengkapnya di Fase_Bulanan)", 17)
        header(ws, start + 1, ["Kode", "Metode", "", "", "", "", "", "", "Puncak"])
        for j in range(min(6, self.n_months)):
            cell = ws.cell(row=start + 1, column=3 + j, value=f"=Fase_Bulanan!{col_letter(3 + j)}5")
            cell.number_format = "mmm-yy"
        row = start + 2
        for method in ("UCP", "Bottom-up"):
            for code in [*roles, "TOTAL"]:
                src = self.fase_rows[(method, code)]
                is_total = code == "TOTAL"
                put(ws, f"A{row}", "Total" if is_total else code, F_BOLD, border=BORDER)
                put(ws, f"B{row}", method, border=BORDER)
                for j in range(min(6, self.n_months)):
                    as_calc(
                        put(
                            ws,
                            f"{col_letter(3 + j)}{row}",
                            f"=Fase_Bulanan!{col_letter(3 + j)}{src}",
                        ),
                        "0.00",
                        RES_FILL if is_total else CALC_FILL,
                        is_total,
                    )
                as_calc(
                    put(ws, f"I{row}", f"=Fase_Bulanan!{self._peak()}{src}"), "0.00", RES_FILL, True
                )
                row += 1
        section(ws, row + 1, "3. Cek kualitas", 17)
        checks = [
            ("Baris breakdown terisi", "=Breakdown_WBS!C3", "0"),
            ("Baris ditandai Perlu review", "=Breakdown_WBS!E3", "0"),
            ("Baris breakdown belum lengkap (harus 0)", "=Breakdown_WBS!H3", "0"),
            (
                "Rating TCF/EF diubah tanpa bukti (harus 0)",
                "=" + "+".join(f'COUNTIF(UCP_{c}!F33:F58,"BUKTI KOSONG")' for c in roles),
                "0",
            ),
            (
                "Cek UUCW vs Breakdown",
                "=IF(AND(" + ",".join(f'UCP_{c}!F27="OK"' for c in roles) + '),"OK","SELISIH")',
                "@",
            ),
            (
                "Cek sebaran bulanan (UCP / Bottom-up)",
                f'=Fase_Bulanan!{self.fase_checks[0]}&" / "&Fase_Bulanan!{self.fase_checks[1]}',
                "@",
            ),
        ]
        for i, (label, formula, fmt) in enumerate(checks):
            r = row + 2 + i
            ws.merge_cells(f"A{r}:E{r}")
            put(ws, f"A{r}", label, border=BORDER)
            as_calc(put(ws, f"F{r}", formula), fmt)
        ws.freeze_panes = "A5"
        page(ws)

    def _peak(self) -> str:
        return col_letter(3 + self.n_months)

    # --------------------------------------------------------------- AI scenario
    def _peluang_ai(self) -> None:
        ws = self.wb.create_sheet("Peluang_AI")
        roles = self.roles
        title(
            ws,
            "PELUANG BANTUAN AI (SKENARIO, BUKAN PENGURANGAN OTOMATIS)",
            "Bagian pekerjaan yang dapat dipercepat dengan bantuan AI. Angka dasar di Ringkasan "
            "tidak berubah. Persentase hemat dan bobot tugas dapat dikoreksi di bagian C.",
            14,
        )
        widths(
            ws,
            {
                "A": 10, "B": 8, "C": 28, "D": 34, "E": 22, "F": 26, "G": 16, "H": 9, "I": 9,
                "J": 10, "K": 26, "L": 40, "M": 11, "N": 11,
            },
        )  # fmt: skip
        n_tasks = max(60, len(self.run.ai_tasks) + 20)
        c_hdr = 10 + len(roles) + 4 + len(self.catalog.ai_categories) + 3
        c0 = c_hdr + 1
        cl = c0 + n_tasks - 1
        section(ws, 4, "A. Ringkasan per role", 14)
        header(
            ws,
            5,
            [
                "Kode", "Role", "PD dasar (bottom-up)", "PD terbantu AI", "Persentase hemat",
                "UCP FTE", "UCP FTE dengan AI", "BU FTE", "BU FTE dengan AI",
            ],
        )  # fmt: skip
        for i, code in enumerate(roles):
            row = 6 + i
            put(ws, f"A{row}", code, F_BOLD, border=BORDER)
            put(ws, f"B{row}", self._name(code), border=BORDER)
            as_calc(
                put(ws, f"C{row}", f"=SUMIFS({self._rng('U')},{self._rng('J')},$A{row})"), "#,##0.0"
            )
            as_calc(
                put(ws, f"D{row}", f"=SUMIFS($N${c0}:$N${cl},$B${c0}:$B${cl},$A{row})"), "#,##0.0"
            )
            as_calc(put(ws, f"E{row}", f"=IF(C{row}=0,0,D{row}/C{row})"), "0%", RES_FILL, True)
            as_calc(put(ws, f"F{row}", f"=Ringkasan!G{7 + i}"), "0.00")
            as_calc(put(ws, f"G{row}", f"=F{row}*(1-E{row})"), "0.00", RES_FILL, True)
            as_calc(put(ws, f"H{row}", f"=Ringkasan!J{7 + i}"), "0.00")
            as_calc(put(ws, f"I{row}", f"=H{row}*(1-E{row})"), "0.00", RES_FILL, True)
        total = 6 + len(roles)
        put(ws, f"B{total}", "Total", F_BOLD, border=BORDER)
        for col in "CDFGHI":
            as_calc(
                put(ws, f"{col}{total}", f"=SUM({col}6:{col}{total - 1})"), "0.00", RES_FILL, True
            )
        as_calc(put(ws, f"E{total}", f"=IF(C{total}=0,0,D{total}/C{total})"), "0%", RES_FILL, True)

        cat_start = total + 2
        section(
            ws, cat_start, "B. Kategori bantuan AI dan rentang hemat (asumsi, perlu kalibrasi)", 14
        )
        header(
            ws,
            cat_start + 1,
            ["Kunci", "", "Kategori", "Contoh", "Hemat minimum", "Hemat maksimum", "Hemat bawaan"],
        )
        for i, cat in enumerate(self.catalog.ai_categories):
            row = cat_start + 2 + i
            put(ws, f"A{row}", cat.key, F_BOLD, border=BORDER)
            put(ws, f"C{row}", cat.label, border=BORDER)
            put(ws, f"D{row}", cat.examples, F_NOTE, align=WRAP)
            for col, value in (
                ("E", cat.min_saving),
                ("F", cat.max_saving),
                ("G", cat.default_saving),
            ):
                as_input(put(ws, f"{col}{row}", value), "0%")

        section(ws, c_hdr - 1, "C. Detail tugas per item", 14)
        header(
            ws,
            c_hdr,
            [
                "WBS ID", "Role", "Item", "Tugas", "Pilar skill", "Skill", "Kategori", "Bobot tugas",
                "Hemat", "Keyakinan", "Prasyarat", "Alasan", "PD baris", "PD terbantu",
            ],
        )  # fmt: skip
        items = {r.key: r for r in self.run.breakdown}
        for k in range(n_tasks):
            row = c0 + k
            task = self.run.ai_tasks[k] if k < len(self.run.ai_tasks) else None
            source_row = items.get(f"{task.item_id}|{task.role}") if task else None
            task_values: list[object] = (
                [
                    task.item_id,
                    task.role,
                    source_row.item_name if source_row else "",
                    task.task,
                    task.pillar,
                    task.skill,
                    task.category,
                    task.weight,
                    task.saving,
                    task.confidence,
                    task.prerequisite,
                    task.reason,
                ]
                if task
                else [None] * 12
            )
            for j, task_value in enumerate(task_values):
                cell = ws.cell(row=row, column=1 + j, value=task_value)
                as_input(cell)
                if j in (7, 8):
                    cell.number_format = "0%"
                if j in (3, 10, 11):
                    cell.alignment = WRAP
            as_calc(
                put(
                    ws,
                    f"M{row}",
                    f'=IF(A{row}="","",SUMIFS({self._rng("U")},{self._rng("B")},A{row},{self._rng("J")},B{row}))',
                ),
                "0.0",
            )
            as_calc(
                put(
                    ws, f"N{row}", f'=IF(OR(M{row}="",H{row}="",I{row}=""),"",M{row}*H{row}*I{row})'
                ),
                "0.00",
            )
        ws.freeze_panes = "A6"
        page(ws)

    # ----------------------------------------------------------------- rubric
    def _rubrik(self) -> None:
        ws = self.wb.create_sheet("Rubrik")
        title(
            ws,
            "RUBRIK DAN KATALOG ROLE",
            "Acuan agent dan reviewer saat mengisi Breakdown_WBS dan UCP.",
            4,
        )
        widths(ws, {"A": 16, "B": 40, "C": 40, "D": 60})
        section(ws, 4, "A. Klasifikasi item", 4)
        header(ws, 5, ["Kategori", "Bobot use case", "", "Ciri"])
        p = self.run.parameters
        for i, (cat, text) in enumerate(self.catalog.classification_rubric.items()):
            put(ws, f"A{6 + i}", cat, F_BOLD, border=BORDER)
            put(ws, f"B{6 + i}", p.use_case_weights.get(cat), border=BORDER)  # type: ignore[call-overload]
            put(ws, f"D{6 + i}", text, align=WRAP, border=BORDER)
        section(ws, 10, "B. Aktor (UAW)", 4)
        header(ws, 11, ["Kategori", "Bobot aktor", "", "Kriteria"])
        for i, (cat, text) in enumerate(self.catalog.actor_rubric.items()):
            put(ws, f"A{12 + i}", cat, F_BOLD, border=BORDER)
            put(ws, f"B{12 + i}", p.actor_weights.get(cat), border=BORDER)  # type: ignore[call-overload]
            put(ws, f"D{12 + i}", text, align=WRAP, border=BORDER)
        section(ws, 16, "C. Aturan keputusan", 4)
        header(ws, 17, ["Kode", "Aturan", "", "Penjelasan"])
        row = 18
        for rule in self.catalog.rules:
            put(ws, f"A{row}", rule.id, F_BOLD, border=BORDER)
            put(ws, f"B{row}", rule.title, border=BORDER)
            put(ws, f"D{row}", rule.text, align=WRAP, border=BORDER)
            ws.row_dimensions[row].height = 40
            row += 1
        section(ws, row + 1, "D. Katalog role", 4)
        header(ws, row + 2, ["Kode", "Lingkup utama", "Di luar lingkup", "Pilar skill"])
        row += 3
        for code in self.roles:
            role = self.catalog.roles.get(code)
            if role is None:
                continue
            put(ws, f"A{row}", f"{code} - {role.name}", F_BOLD, border=BORDER, align=WRAP)
            put(ws, f"B{row}", role.in_scope, align=WRAP, border=BORDER)
            put(ws, f"C{row}", role.out_of_scope, align=WRAP, border=BORDER)
            put(ws, f"D{row}", "; ".join(role.pillars), align=WRAP, border=BORDER)
            ws.row_dimensions[row].height = 70
            row += 1
        page(ws)

    # -------------------------------------------------------------- capability
    def _capability(self) -> None:
        ws = self.wb.create_sheet("Capability")
        title(
            ws,
            "CAPABILITY PER ITEM WBS",
            "Satu baris = satu skill yang dibutuhkan satu role pada satu item WBS. Satu item "
            "bisa punya beberapa baris. Pakai filter atau pivot untuk mengolahnya.",
            9,
        )
        widths(
            ws,
            {"A": 5, "B": 10, "C": 32, "D": 16, "E": 8, "F": 20, "G": 26, "H": 32, "I": 40},
        )
        header(
            ws,
            4,
            ["No", "WBS ID", "Item WBS", "Kelompok", "Role", "Nama role", "Pilar", "Skill",
             "Bagian pekerjaan"],
        )  # fmt: skip
        ws.row_dimensions[4].height = 30
        detail = [(row, cap) for row in self.run.breakdown for cap in row.capabilities]
        first = 5
        n = max(40, len(detail) + 20)
        for i in range(n):
            r = first + i
            put(ws, f"A{r}", i + 1, align=CENTER, border=BORDER)
            entry = detail[i] if i < len(detail) else None
            values = (
                (entry[0].item_id, entry[0].item_name, entry[0].group, entry[0].role,
                 self._name(entry[0].role), entry[1].pillar, entry[1].skill, entry[1].reason)
                if entry else (None,) * 8
            )  # fmt: skip
            for col, value in zip("BCDEFGHI", values, strict=True):
                as_input(put(ws, f"{col}{r}", value))
                if col in "CHI":
                    ws[f"{col}{r}"].alignment = WRAP
        ws.auto_filter.ref = f"A4:I{first + n - 1}"
        ws.freeze_panes = "A5"
        page(ws)

    # --------------------------------------------------------------- findings
    def _asumsi(self) -> None:
        ws = self.wb.create_sheet("Asumsi_Log")
        title(
            ws,
            "LOG ASUMSI, TEMUAN, DAN KEPUTUSAN",
            "Temuan agent dan asumsi hitung. Ubah kolom Status setelah direview.",
            6,
        )
        widths(ws, {"A": 5, "B": 18, "C": 80, "D": 26, "E": 12, "F": 16})
        header(ws, 4, ["No", "Area", "Asumsi / temuan", "Sumber", "Dampak", "Status"])
        p = self.run.parameters
        entries = [
            (
                "Metode",
                "UCP diadaptasi ke item WBS. Bobot 5/10/15 dan PHM belum dikalibrasi untuk "
                "semua jenis proyek; gunakan sebagai pembanding.",
                "Karner 1993",
                "Tinggi",
            ),
            (
                "Parameter",
                f"PHM {p.phm:g} jam/UCP; {p.days_per_month:g} hari kerja/bulan; "
                f"{p.hours_per_day:g} jam/hari; faktor pendukung {p.support_factor:g}; "
                f"toleransi {p.tolerance:.0%}.",
                "Parameter",
                "Tinggi",
            ),
        ]
        entries += [
            (
                FINDING_LABELS.get(f.kind, f.kind),
                (f"[{f.item_id}{' / ' + f.role if f.role else ''}] " if f.item_id else "")
                + f.message,
                f.source,
                f.severity,
            )
            for f in self.run.findings
        ]
        n = max(60, len(entries) + 20)
        for i in range(n):
            row = 5 + i
            put(ws, f"A{row}", i + 1, align=CENTER, border=BORDER)
            entry = entries[i] if i < len(entries) else None
            for col, value in zip("BCDE", entry or (None,) * 4, strict=True):
                as_input(put(ws, f"{col}{row}", value))
                ws[f"{col}{row}"].alignment = WRAP
            as_input(put(ws, f"F{row}", "Perlu review" if entry else None))
        dv = DataValidation(
            type="list", formula1='"Perlu review,Disetujui,Ditolak,Selesai"', allow_blank=True
        )
        ws.add_data_validation(dv)
        dv.add(f"F5:F{4 + n}")
        ws.freeze_panes = "A5"
        page(ws)

    # --------------------------------------------------------------- guidance
    def _petunjuk(self) -> None:
        ws = self.wb["Sheet"]
        ws.title = "Petunjuk"
        title(
            ws,
            "VALIDASI KEBUTUHAN TALENT",
            "Dua pendekatan: (1) Use Case Points dan (2) Bottom-Up dari jadwal WBS. "
            "Hasil agent; perlu review sebelum dipakai.",
            8,
        )
        widths(ws, {"A": 4, "B": 24, "C": 18, "D": 16, "E": 11, "F": 11, "G": 9, "H": 60})
        section(ws, 4, "Cara membaca", 8)
        steps = [
            "1. Ringkasan: kebutuhan per role dari dua metode, dibandingkan dengan permintaan.",
            "2. Breakdown_WBS: asal setiap angka. Koreksi baris yang ditandai Perlu review; "
            "semua total ikut berubah.",
            "3. Capability: skill yang dibutuhkan tiap item WBS per role, siap difilter.",
            "4. UCP per role: aktor dan rating faktor. Rating selain 3 wajib punya bukti.",
            "5. Fase_Bulanan: sebaran kebutuhan per bulan dan puncaknya.",
            "6. Peluang_AI: skenario bantuan AI per tugas dan capability.",
            "7. Asumsi_Log: temuan agent dan asumsi hitung.",
        ]
        for i, text in enumerate(steps):
            ws.merge_cells(f"A{5 + i}:H{5 + i}")
            put(ws, f"A{5 + i}", text, align=WRAP)
            ws.row_dimensions[5 + i].height = 24
        section(ws, 12, "Legenda warna", 8)
        legend = [
            ("Sel input", F_IN, IN_FILL, "Kuning, tulisan biru: diisi agent atau reviewer."),
            (
                "Sel rumus",
                F_NORM,
                CALC_FILL,
                "Abu-abu: jangan ditimpa. Semua perhitungan dilakukan rumus.",
            ),
            ("Hasil kunci", F_BOLD, RES_FILL, "Hijau: angka utama."),
            ("Perlu review", F_NORM, WARN_FILL, "Oranye: butuh keputusan reviewer."),
        ]
        for i, (label, font, fill, text) in enumerate(legend):
            row = 13 + i
            put(ws, f"B{row}", label, font, fill, border=BORDER)
            ws.merge_cells(f"C{row}:H{row}")
            put(ws, f"C{row}", text)
        page(ws)

    # ------------------------------------------------------------------ public
    def build(self) -> bytes:
        """Create the workbook and return it as bytes.

        Returns
        -------
        bytes
            Content of an ``.xlsx`` file.
        """
        self._parameter()
        self._breakdown()
        for code in self.roles:
            self._ucp(code)
        self._capability()
        self._bottom_up()
        self._fase()
        self._ringkasan()
        if self.run.ai_enabled:
            self._peluang_ai()
        self._rubrik()
        self._asumsi()
        self._petunjuk()
        order = [
            "Petunjuk", "Parameter", "Ringkasan", "Breakdown_WBS", "Capability",
            *[f"UCP_{c}" for c in self.roles], "BottomUp", "Fase_Bulanan",
            *(["Peluang_AI"] if self.run.ai_enabled else []), "Rubrik", "Asumsi_Log",
        ]  # fmt: skip
        self.wb._sheets = [self.wb[name] for name in order]
        colors = {
            "Petunjuk": "7F7F7F", "Parameter": "FFC000", "Ringkasan": "70AD47",
            "Breakdown_WBS": "FFC000", "Capability": "FFC000", "Peluang_AI": "7C3AED", "Rubrik": "7F7F7F",
            "Asumsi_Log": "FFC000",
        }  # fmt: skip
        for name in order:
            self.wb[name].sheet_properties.tabColor = colors.get(name, "5B9BD5")
        self.wb.active = 2
        self.wb.calculation.fullCalcOnLoad = True
        buffer = io.BytesIO()
        self.wb.save(buffer)
        return buffer.getvalue()


def build_workbook(run: RunResult, catalog: Catalog) -> bytes:
    """Build the validation workbook for a run.

    Parameters
    ----------
    run : RunResult
        A recalculated run.
    catalog : Catalog
        Capability catalog used in the run.

    Returns
    -------
    bytes
        The ``.xlsx`` file content.

    Notes
    -----
    The workbook stores formulas, not values. Excel computes them when the
    file is opened.
    """
    return WorkbookBuilder(run, catalog).build()
