"""Cell styles and small helpers used by the workbook builder."""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="2E1A47")
SEC_FILL = PatternFill("solid", fgColor="E9E3F5")
IN_FILL = PatternFill("solid", fgColor="FFF9DB")
CALC_FILL = PatternFill("solid", fgColor="F2F2F2")
RES_FILL = PatternFill("solid", fgColor="E2F0D9")
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
_THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

F_NORM = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True, color="2E1A47")
F_SUB = Font(name=FONT, size=10, italic=True, color="595959")
F_IN = Font(name=FONT, size=10, color="0000FF")
F_NOTE = Font(name=FONT, size=9, italic=True, color="595959")

WRAP = Alignment(wrap_text=True, vertical="top")
WRAPC = Alignment(wrap_text=True, vertical="center", horizontal="center")
CENTER = Alignment(horizontal="center", vertical="center")


def put(
    ws: Worksheet,
    addr: str,
    value: object,
    font: Font = F_NORM,
    fill: PatternFill | None = None,
    align: Alignment | None = None,
    border: Border | None = None,
    fmt: str | None = None,
) -> object:
    """Write a value and style into one cell and return the cell."""
    cell = ws[addr]
    cell.value = value
    cell.font = font
    if fill:
        cell.fill = fill
    if align:
        cell.alignment = align
    if border:
        cell.border = border
    if fmt:
        cell.number_format = fmt
    return cell


def as_input(cell: object, fmt: str | None = None) -> None:
    """Style a cell as user input (yellow fill, blue text)."""
    cell.font = F_IN  # type: ignore[attr-defined]
    cell.fill = IN_FILL  # type: ignore[attr-defined]
    cell.border = BORDER  # type: ignore[attr-defined]
    if fmt:
        cell.number_format = fmt  # type: ignore[attr-defined]


def as_calc(
    cell: object, fmt: str | None = None, fill: PatternFill = CALC_FILL, bold: bool = False
) -> None:
    """Style a cell as a formula (grey fill)."""
    cell.font = F_BOLD if bold else F_NORM  # type: ignore[attr-defined]
    cell.fill = fill  # type: ignore[attr-defined]
    cell.border = BORDER  # type: ignore[attr-defined]
    if fmt:
        cell.number_format = fmt  # type: ignore[attr-defined]


def title(ws: Worksheet, text: str, sub: str, width_cols: int) -> None:
    """Write the sheet title and subtitle on rows 1 and 2."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width_cols)
    put(ws, "A1", text, F_TITLE)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width_cols)
    put(ws, "A2", sub, F_SUB, align=WRAP)
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 28


def section(ws: Worksheet, row: int, text: str, ncols: int) -> None:
    """Write a shaded section header spanning ``ncols`` columns."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    for col in range(1, ncols + 1):
        ws.cell(row=row, column=col).fill = SEC_FILL
    put(ws, f"A{row}", text, F_BOLD, SEC_FILL)


def header(ws: Worksheet, row: int, labels: list[str], start_col: int = 1) -> None:
    """Write a dark table header row."""
    for i, text in enumerate(labels):
        cell = ws.cell(row=row, column=start_col + i, value=text)
        cell.font = F_HDR
        cell.fill = HDR_FILL
        cell.alignment = WRAPC
        cell.border = BORDER


def widths(ws: Worksheet, mapping: dict[str, float]) -> None:
    """Set column widths."""
    for col, width in mapping.items():
        ws.column_dimensions[col].width = width


def page(ws: Worksheet, landscape: bool = True) -> None:
    """Fit the sheet to one page wide when printed."""
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
