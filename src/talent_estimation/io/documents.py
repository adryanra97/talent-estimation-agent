"""Extract plain text from uploaded documents, in memory only."""

from __future__ import annotations

import io

MAX_CHARS = 60_000
"""Upper bound of text passed on to a language model per document."""


def extract_text(filename: str, content: bytes) -> str:
    """Return the text content of a PDF, Word, Markdown, or plain-text file.

    Parameters
    ----------
    filename : str
        Original file name, used to detect the format.
    content : bytes
        File content.

    Returns
    -------
    str
        Extracted text, truncated to :data:`MAX_CHARS` characters.

    Raises
    ------
    ValueError
        If the format is not supported.
    """
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif name.endswith(".docx"):
        from docx import Document

        document = Document(io.BytesIO(content))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts)
    elif name.endswith((".txt", ".md")):
        text = content.decode("utf-8", errors="replace")
    else:
        raise ValueError("Dokumen harus berformat .pdf, .docx, .txt, atau .md.")
    return text[:MAX_CHARS]
