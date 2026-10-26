"""Read a work breakdown structure from an Excel workbook.

The reader looks for a header row with recognisable column names, keeps only
leaf work packages, and collects a legend of role codes when the workbook has
a team sheet. Person names in a team sheet are never returned.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook
from pydantic import BaseModel

from talent_estimation.calc.calendar import networkdays
from talent_estimation.domain.models import WbsItem

COLUMN_KEYWORDS: dict[str, tuple[str, ...]] = {
    "description": ("deskripsi", "description", "deliverable", "detail", "uraian"),
    "role": ("talent", "role", "pic", "resource", "assignee", "penanggung"),
    "start": ("start", "mulai", "begin"),
    "end": ("end", "finish", "selesai", "akhir"),
    "workdays": ("wd", "working day", "hari kerja", "duration", "durasi", "days"),
    "item_id": ("wbs", "id", "kode", "no"),
    "name": (
        "work package",
        "workstream",
        "item",
        "task",
        "aktivitas",
        "activity",
        "nama",
        "pekerjaan",
    ),
}
"""Keywords per field, matched against header cells in this order."""

REQUIRED = ("name", "start", "end")


class ColumnMapping(BaseModel):
    """Zero-based column indexes of a WBS table."""

    sheet: str
    header_row: int
    item_id: int | None = None
    name: int
    description: int | None = None
    role: int | None = None
    start: int
    end: int
    workdays: int | None = None


@dataclass
class WbsReadResult:
    """Items and side information read from a WBS workbook."""

    items: list[WbsItem]
    mapping: ColumnMapping
    role_legend: dict[str, str] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    total_rows: int = 0


class WbsFormatError(ValueError):
    """Raised when no WBS table can be recognised in a workbook."""


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d %b %Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                continue
    return None


def _normalise_id(value: Any) -> str:
    text = str(value).strip()
    text = re.sub(r"(\.0)+$", "", text)
    return text


def _match_header(row: tuple[Any, ...], sheet: str, index: int) -> ColumnMapping | None:
    cells = [str(c).strip().lower() if c is not None else "" for c in row]
    found: dict[str, int] = {}
    for fieldname, keywords in COLUMN_KEYWORDS.items():
        for col, text in enumerate(cells):
            if not text or col in found.values():
                continue
            if fieldname in ("item_id",):
                if text in keywords or text.startswith(("wbs", "id ")):
                    found[fieldname] = col
                    break
            elif any(k in text for k in keywords):
                found[fieldname] = col
                break
    if not all(k in found for k in REQUIRED):
        return None
    return ColumnMapping(sheet=sheet, header_row=index, **found)


def detect_columns(content: bytes) -> ColumnMapping:
    """Find the WBS table in a workbook.

    Parameters
    ----------
    content : bytes
        Excel file content.

    Returns
    -------
    ColumnMapping
        The detected sheet, header row, and columns.

    Raises
    ------
    WbsFormatError
        If no sheet has a header row with name, start, and end columns.
    """
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        for index, row in enumerate(sheet.iter_rows(min_row=1, max_row=30, values_only=True)):
            mapping = _match_header(row, sheet.title, index)
            if mapping:
                return mapping
    raise WbsFormatError(
        "Tabel WBS tidak dikenali. Pastikan ada baris judul kolom dengan nama item, "
        "tanggal mulai, dan tanggal selesai."
    )


def preview_rows(content: bytes, limit: int = 40) -> dict[str, list[list[str]]]:
    """Return the first rows of every sheet as text, for column inference by a model."""
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    result: dict[str, list[list[str]]] = {}
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(min_row=1, max_row=limit, values_only=True):
            rows.append(["" if c is None else str(c)[:60] for c in row[:15]])
        result[sheet.title] = rows
    return result


def _role_legend(workbook: Any, tokens: set[str], wbs_sheet: str) -> dict[str, str]:
    """Map role codes used in the WBS to role titles found in another sheet.

    Only the code and the role title are kept. Other columns of the team sheet,
    such as person names, are ignored.
    """
    legend: dict[str, str] = {}
    for sheet in workbook.worksheets:
        if sheet.title == wbs_sheet:
            continue
        header_role_col: int | None = None
        for row in sheet.iter_rows(min_row=1, max_row=300, values_only=True):
            cells = ["" if c is None else str(c).strip() for c in row]
            if header_role_col is None:
                lowered = [c.lower() for c in cells]
                if "role" in lowered:
                    header_role_col = lowered.index("role")
                    continue
            if header_role_col is None:
                continue
            for col, value in enumerate(cells):
                if value in tokens and header_role_col < len(cells) and col != header_role_col:
                    title = cells[header_role_col]
                    if title and value not in legend:
                        legend[value] = title
    return legend


def read_wbs(
    content: bytes,
    mapping: ColumnMapping | None = None,
    holidays: list[date] | None = None,
) -> WbsReadResult:
    """Read leaf work packages from a WBS workbook.

    Parameters
    ----------
    content : bytes
        Excel file content.
    mapping : ColumnMapping, optional
        Column mapping. Detected automatically when omitted.
    holidays : list of date, optional
        Holidays used when working days must be computed from dates.

    Returns
    -------
    WbsReadResult
        Leaf items, the mapping used, a role legend, and warnings.
    """
    mapping = mapping or detect_columns(content)
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook[mapping.sheet]
    rows = list(sheet.iter_rows(min_row=mapping.header_row + 2, values_only=True))

    def cell(row: tuple[Any, ...], index: int | None) -> Any:
        if index is None or index >= len(row):
            return None
        return row[index]

    raw: list[dict[str, Any]] = []
    for row in rows:
        name = cell(row, mapping.name)
        if name is None or str(name).strip() == "":
            continue
        raw.append(
            {
                "id": _normalise_id(cell(row, mapping.item_id))
                if cell(row, mapping.item_id) is not None
                else "",
                "name": str(name).strip(),
                "description": str(cell(row, mapping.description) or "").strip(),
                "role": str(cell(row, mapping.role) or "").strip(),
                "start": _as_date(cell(row, mapping.start)),
                "end": _as_date(cell(row, mapping.end)),
                "workdays": cell(row, mapping.workdays),
            }
        )

    ids = [r["id"] for r in raw if r["id"]]
    names_by_id = {r["id"]: r["name"] for r in raw if r["id"]}
    warnings: list[str] = []
    items: list[WbsItem] = []
    for index, r in enumerate(raw, start=1):
        item_id = r["id"] or f"ROW-{index}"
        if r["id"] and any(other.startswith(r["id"] + ".") for other in ids):
            continue
        workdays = r["workdays"]
        if not isinstance(workdays, int | float):
            workdays = None
        if workdays is None and r["start"] and r["end"]:
            workdays = networkdays(r["start"], r["end"], holidays or [])
        if r["start"] is None or r["end"] is None:
            warnings.append(f"Item {item_id} tidak punya tanggal mulai atau selesai.")
        parent = item_id.rsplit(".", 1)[0] if "." in item_id else ""
        items.append(
            WbsItem(
                item_id=item_id,
                name=r["name"],
                description=r["description"],
                role_hint=r["role"],
                start=r["start"],
                end=r["end"],
                workdays=float(workdays) if workdays is not None else None,
                group=names_by_id.get(parent, ""),
            )
        )

    tokens = {
        token.strip()
        for item in items
        for token in re.split(r"[,;/]", item.role_hint)
        if token.strip()
    }
    legend = _role_legend(workbook, tokens, mapping.sheet)
    return WbsReadResult(
        items=items, mapping=mapping, role_legend=legend, warnings=warnings, total_rows=len(raw)
    )


def _file_label(name: str) -> str:
    stem = re.sub(r"\.[^.]+$", "", name or "").strip()
    return stem or "WBS"


def merge_wbs(results: list[tuple[str, WbsReadResult]]) -> WbsReadResult:
    """Combine several WBS workbooks into one list of items.

    Parameters
    ----------
    results : list of (str, WbsReadResult)
        File name and read result per workbook, in upload order.

    Returns
    -------
    WbsReadResult
        All leaf items. When item IDs repeat across files, every item gets a
        prefix ``A:``, ``B:`` and so on by file so that each row stays unique.
        Items without a group use the file name as group. The mapping is the
        one of the first file.
    """
    if not results:
        raise WbsFormatError("Tidak ada file WBS.")
    if len(results) == 1:
        return results[0][1]
    seen: set[str] = set()
    clash = False
    for _, result in results:
        ids = {i.item_id for i in result.items}
        clash = clash or bool(ids & seen)
        seen |= ids
    items: list[WbsItem] = []
    legend: dict[str, str] = {}
    warnings: list[str] = []
    for index, (name, result) in enumerate(results):
        label = _file_label(name)
        prefix = f"{chr(65 + index)}:" if clash else ""
        for item in result.items:
            items.append(
                item.model_copy(
                    update={
                        "item_id": f"{prefix}{item.item_id}",
                        "group": item.group or label,
                    }
                )
            )
        legend.update(result.role_legend)
        warnings.extend(f"{label}: {w}" for w in result.warnings)
    if clash:
        names = ", ".join(f"{chr(65 + i)} = {_file_label(n)}" for i, (n, _) in enumerate(results))
        warnings.append(f"ID item sama muncul di lebih dari satu file. ID diberi awalan: {names}.")
    return WbsReadResult(
        items=items,
        mapping=results[0][1].mapping,
        role_legend=legend,
        warnings=warnings,
        total_rows=sum(r.total_rows for _, r in results),
    )
