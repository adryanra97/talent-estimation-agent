import io
import shutil
import subprocess

import pytest
from openpyxl import load_workbook

from talent_estimation.agents import run_pipeline
from talent_estimation.domain.models import RunResult
from talent_estimation.excel import build_workbook


async def demo_run(inputs):
    async for e in run_pipeline(inputs, None):
        if e["type"] == "result":
            return RunResult.model_validate(e["run"])


async def test_workbook_structure(demo_inputs, catalog):
    run = await demo_run(demo_inputs)
    wb = load_workbook(io.BytesIO(build_workbook(run, catalog)))
    assert wb.sheetnames[:4] == ["Petunjuk", "Parameter", "Ringkasan", "Breakdown_WBS"]
    assert {"UCP_DE", "UCP_AI", "Peluang_AI", "Asumsi_Log"} <= set(wb.sheetnames)
    assert wb["Breakdown_WBS"]["B7"].value == run.breakdown[0].item_id


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
async def test_excel_formulas_match_engine(demo_inputs, catalog, tmp_path):
    run = await demo_run(demo_inputs)
    src = tmp_path / "run.xlsx"
    src.write_bytes(build_workbook(run, catalog))
    out = tmp_path / "out"
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
        check=True,
        capture_output=True,
        timeout=180,
    )
    sheet = load_workbook(out / "run.xlsx", data_only=True)["Ringkasan"]
    for i, e in enumerate(run.estimates):
        r = 7 + i
        assert sheet[f"G{r}"].value == pytest.approx(e.ucp_fte, abs=1e-6)
        assert sheet[f"J{r}"].value == pytest.approx(e.bu_fte, abs=1e-6)
        assert sheet[f"H{r}"].value == pytest.approx(e.ucp_peak_fte, abs=1e-6)
        assert sheet[f"K{r}"].value == pytest.approx(e.bu_peak_fte, abs=1e-6)
