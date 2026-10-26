import io

import pytest
from openpyxl import Workbook

from talent_estimation.io import WbsFormatError, extract_text, read_wbs
from talent_estimation.llm import mask_text


def test_demo_wbs_leaf_items(demo_wbs):
    result = read_wbs(demo_wbs)
    ids = [i.item_id for i in result.items]
    assert "2.1.2" in ids and "2.1" not in ids and "1" not in ids
    assert all(i.workdays for i in result.items)
    assert next(i for i in result.items if i.item_id == "2.1.2").group == "Pengaduan Warga"


def test_role_legend_ignores_other_columns():
    wb = Workbook()
    ws = wb.active
    ws.append(["WBS", "Task", "Talent", "Start", "End"])
    ws.append(["1.1", "Integrasi", "DE1", "2027-02-01", "2027-02-10"])
    team = wb.create_sheet("Team")
    team.append(["Code", "Role", "Nama"])
    team.append(["DE1", "Data Engineer", "Orang Rahasia"])
    buf = io.BytesIO()
    wb.save(buf)
    result = read_wbs(buf.getvalue())
    assert result.role_legend == {"DE1": "Data Engineer"}
    assert "Rahasia" not in str(result.role_legend)


def test_unrecognised_workbook():
    wb = Workbook()
    wb.active.append(["a", "b"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(WbsFormatError):
        read_wbs(buf.getvalue())


def test_text_and_masking():
    assert extract_text("a.md", b"halo") == "halo"
    with pytest.raises(ValueError):
        extract_text("a.exe", b"")
    masked = mask_text("Kirim ke a.b@contoh.id, 081234567890, NIK 3201234567890123")
    assert "[EMAIL]" in masked and "[TELEPON]" in masked and "[NIK]" in masked
